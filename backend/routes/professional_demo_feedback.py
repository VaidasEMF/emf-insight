from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from typing import Optional
from sqlalchemy import text
from sqlalchemy.orm import Session
from docx import Document

from db.database import SessionLocal
from models.user import User
from routes.auth import get_current_user


router = APIRouter(
    prefix="/professional-demo-feedback",
    tags=["professional-demo-feedback"],
)


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


class DemoFeedbackRequest(BaseModel):
    overall_rating: int | None = None
    business_survey_rating: int | None = None
    measurements_rating: int | None = None
    results_rating: int | None = None
    pdf_rating: int | None = None
    experience_level: Optional[str] = None
    experience_years: Optional[str] = None
    background_text: Optional[str] = None
    home_projects_clarity: Optional[str] = None
    home_projects_clarity_text: Optional[str] = None
    work_type: Optional[str] = None

    confusing_text: str | None = None
    missing_features: str | None = None
    bugs_text: str | None = None
    improvements_text: str | None = None

    professional_use: bool | None = None
    recommendation_score: int | None = None


def validate_feedback(payload: DemoFeedbackRequest):
    ratings = {
        "overall_rating": payload.overall_rating,
        "business_survey_rating": payload.business_survey_rating,
        "measurements_rating": payload.measurements_rating,
        "results_rating": payload.results_rating,
        "pdf_rating": payload.pdf_rating,
    }

    for field, value in ratings.items():
        if value is not None and not 1 <= value <= 5:
            raise HTTPException(
                status_code=400,
                detail=f"{field} must be between 1 and 5",
            )

    if (
        payload.recommendation_score is not None
        and not 0 <= payload.recommendation_score <= 10
    ):
        raise HTTPException(
            status_code=400,
            detail="recommendation_score must be between 0 and 10",
        )


@router.post("")
def submit_demo_feedback(
    payload: DemoFeedbackRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    if current_user.role != "professional":
        raise HTTPException(
            status_code=403,
            detail="Professional account required",
        )

    validate_feedback(payload)

    demo = db.execute(
        text(
            """
            SELECT
                id,
                status,
                user_id
            FROM professional_demo_links
            WHERE user_id = :user_id
            ORDER BY activated_at DESC NULLS LAST
            LIMIT 1
            """
        ),
        {"user_id": str(current_user.id)},
    ).mappings().first()

    if not demo:
        raise HTTPException(
            status_code=404,
            detail="Professional Demo not found",
        )

    existing = db.execute(
        text(
            """
            SELECT id
            FROM professional_demo_feedback
            WHERE demo_id = :demo_id
            LIMIT 1
            """
        ),
        {"demo_id": demo["id"]},
    ).mappings().first()

    if existing:
        raise HTTPException(
            status_code=409,
            detail="Feedback already submitted",
        )

    db.execute(
        text(
            """
            INSERT INTO professional_demo_feedback (
                demo_id,
                user_id,
                overall_rating,
                business_survey_rating,
                measurements_rating,
                results_rating,
                pdf_rating,
                experience_level,
                experience_years,
                background_text,
                home_projects_clarity,
                home_projects_clarity_text,
                work_type,
                confusing_text,
                missing_features,
                bugs_text,
                improvements_text,
                professional_use,
                recommendation_score
            )
            VALUES (
                :demo_id,
                :user_id,
                :overall_rating,
                :business_survey_rating,
                :measurements_rating,
                :results_rating,
                :pdf_rating,
                :experience_level,
                :experience_years,
                :background_text,
                :home_projects_clarity,
                :home_projects_clarity_text,
                :work_type,
                :confusing_text,
                :missing_features,
                :bugs_text,
                :improvements_text,
                :professional_use,
                :recommendation_score
            )
            """
        ),
        {
            "demo_id": demo["id"],
            "user_id": str(current_user.id),
            "overall_rating": payload.overall_rating,
            "business_survey_rating": payload.business_survey_rating,
            "measurements_rating": payload.measurements_rating,
            "results_rating": payload.results_rating,
            "pdf_rating": payload.pdf_rating,
            "experience_level": payload.experience_level,
            "experience_years": payload.experience_years,
            "background_text": payload.background_text,
            "home_projects_clarity": payload.home_projects_clarity,
            "home_projects_clarity_text": payload.home_projects_clarity_text,
            "work_type": payload.work_type,
            "confusing_text": payload.confusing_text,
            "missing_features": payload.missing_features,
            "bugs_text": payload.bugs_text,
            "improvements_text": payload.improvements_text,
            "professional_use": payload.professional_use,
            "recommendation_score": payload.recommendation_score,
        },
    )

    db.commit()

    return {
        "status": "received",
        "message": "Demo feedback submitted",
    }


@router.get("/my")
def get_my_demo_feedback(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    if current_user.role != "professional":
        raise HTTPException(
            status_code=403,
            detail="Professional account required",
        )

    feedback = db.execute(
        text(
            """
            SELECT
                f.id,
                f.demo_id,
                f.user_id,
                f.overall_rating,
                f.business_survey_rating,
                f.measurements_rating,
                f.results_rating,
                f.pdf_rating,
                f.experience_level,
                f.experience_years,
                f.background_text,
                f.home_projects_clarity,
                f.home_projects_clarity_text,
                f.work_type,
                f.confusing_text,
                f.missing_features,
                f.bugs_text,
                f.improvements_text,
                f.professional_use,
                f.recommendation_score,
                f.created_at,
                f.updated_at
            FROM professional_demo_feedback f
            WHERE f.user_id = :user_id
            ORDER BY f.created_at DESC
            LIMIT 1
            """
        ),
        {"user_id": str(current_user.id)},
    ).mappings().first()

    if not feedback:
        return {
            "submitted": False,
            "feedback": None,
        }

    return {
        "submitted": True,
        "feedback": dict(feedback),
    }


@router.patch("/my")
def update_my_demo_feedback(
    payload: DemoFeedbackRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    if current_user.role != "professional":
        raise HTTPException(
            status_code=403,
            detail="Professional account required",
        )

    validate_feedback(payload)

    feedback = db.execute(
        text(
            """
            SELECT id
            FROM professional_demo_feedback
            WHERE user_id = :user_id
            ORDER BY created_at DESC
            LIMIT 1
            """
        ),
        {"user_id": str(current_user.id)},
    ).mappings().first()

    if not feedback:
        raise HTTPException(
            status_code=404,
            detail="Feedback not found",
        )

    fields = payload.model_dump()

    db.execute(
        text(
            """
            UPDATE professional_demo_feedback
            SET
                overall_rating = :overall_rating,
                business_survey_rating = :business_survey_rating,
                measurements_rating = :measurements_rating,
                results_rating = :results_rating,
                pdf_rating = :pdf_rating,
                confusing_text = :confusing_text,
                missing_features = :missing_features,
                bugs_text = :bugs_text,
                improvements_text = :improvements_text,
                professional_use = :professional_use,
                recommendation_score = :recommendation_score,
                updated_at = CURRENT_TIMESTAMP
            WHERE id = :id
            """
        ),
        {
            **fields,
            "id": feedback["id"],
        },
    )

    db.commit()

    return {
        "status": "updated",
        "message": "Demo feedback updated",
    }


@router.get("/admin/{demo_id}")
def get_demo_feedback_admin(
    demo_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    if not current_user.is_admin:
        raise HTTPException(
            status_code=403,
            detail="Admin access required",
        )

    feedback = db.execute(
        text(
            """
            SELECT
                f.id,
                f.demo_id,
                f.user_id,

                u.first_name,
                u.last_name,
                u.email,

                p.country,
                p.city,

                f.overall_rating,
                f.business_survey_rating,
                f.measurements_rating,
                f.results_rating,
                f.pdf_rating,

                f.experience_level,
                f.experience_years,
                f.background_text,
                f.home_projects_clarity,
                f.home_projects_clarity_text,
                f.work_type,

                f.confusing_text,
                f.missing_features,
                f.bugs_text,
                f.improvements_text,

                f.professional_use,
                f.recommendation_score,

                f.created_at,
                f.updated_at

            FROM professional_demo_feedback f

            LEFT JOIN users u
                ON CAST(u.id AS VARCHAR) = f.user_id

            LEFT JOIN professional_profiles p
                ON CAST(p.user_id AS VARCHAR) = f.user_id

            WHERE f.demo_id = :demo_id
            LIMIT 1
            """
        ),
        {"demo_id": demo_id},
        ).mappings().first()

    if not feedback:
        return {
            "submitted": False,
            "demo_id": demo_id,
        }

    return {
        "submitted": True,
        "feedback": dict(feedback),
    }


@router.get("/admin/{demo_id}/docx")
def export_demo_feedback_docx(
    demo_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    if not current_user.is_admin:
        raise HTTPException(
            status_code=403,
            detail="Admin access required",
        )

    feedback = db.execute(
        text(
            """
            SELECT
                f.demo_id,
                u.first_name,
                u.last_name,
                u.email,
                p.country,
                p.city,

                f.overall_rating,
                f.business_survey_rating,
                f.measurements_rating,
                f.results_rating,
                f.pdf_rating,

                f.experience_level,
                f.experience_years,
                f.background_text,
                f.home_projects_clarity,
                f.home_projects_clarity_text,
                f.work_type,

                f.professional_use,
                f.recommendation_score,

                f.confusing_text,
                f.missing_features,
                f.bugs_text,
                f.improvements_text,

                f.created_at

            FROM professional_demo_feedback f

            LEFT JOIN users u
                ON CAST(u.id AS VARCHAR) = f.user_id

            LEFT JOIN professional_profiles p
                ON CAST(p.user_id AS VARCHAR) = f.user_id

            WHERE f.demo_id = :demo_id
            LIMIT 1
            """
        ),
        {"demo_id": demo_id},
    ).mappings().first()

    if not feedback:
        raise HTTPException(
            status_code=404,
            detail="Feedback not found",
        )

    def display(value):
        if value is None or value == "":
            return "—"
        return str(value)

    document = Document()

    document.add_heading(
        "EMF Insight — Professional Demo Feedback",
        level=0,
    )

    document.add_paragraph(
        "Tester feedback submitted for Professional Demo."
    )

    document.add_heading("Tester Information", level=1)

    tester_fields = [
        ("First Name", feedback["first_name"]),
        ("Last Name", feedback["last_name"]),
        ("Email", feedback["email"]),
        ("Country", feedback["country"]),
        ("City", feedback["city"]),
        ("Feedback Date", feedback["created_at"]),
    ]

    table = document.add_table(rows=0, cols=2)
    table.style = "Light Shading Accent 1"

    for label, value in tester_fields:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = display(value)

    sections = [
        (
            "Product Experience",
            [
                ("Overall Rating", feedback["overall_rating"]),
                ("Business Survey", feedback["business_survey_rating"]),
                ("Measurements", feedback["measurements_rating"]),
                ("Results", feedback["results_rating"]),
                ("PDF", feedback["pdf_rating"]),
            ],
        ),
        (
            "Professional Background",
            [
                ("EMF Experience", feedback["experience_level"]),
                ("Years of Experience", feedback["experience_years"]),
                ("Main Type of Work", feedback["work_type"]),
                ("Background", feedback["background_text"]),
            ],
        ),
        (
            "Home Projects",
            [
                ("Home Projects Clarity", feedback["home_projects_clarity"]),
                (
                    "What Would Make Home Projects Clearer?",
                    feedback["home_projects_clarity_text"],
                ),
            ],
        ),
        (
            "Professional Use",
            [
                ("Professional Use", feedback["professional_use"]),
                (
                    "Recommendation Score",
                    feedback["recommendation_score"],
                ),
            ],
        ),
        (
            "Open Feedback",
            [
                ("What was confusing?", feedback["confusing_text"]),
                ("Missing Features", feedback["missing_features"]),
                ("Bugs", feedback["bugs_text"]),
                ("Improvements", feedback["improvements_text"]),
            ],
        ),
    ]

    for section_title, fields in sections:
        document.add_heading(section_title, level=1)

        table = document.add_table(rows=0, cols=2)
        table.style = "Light Shading Accent 1"

        for label, value in fields:
            cells = table.add_row().cells
            cells[0].text = label
            cells[1].text = display(value)

    output = BytesIO()
    document.save(output)
    output.seek(0)

    filename = f"professional_demo_feedback_{demo_id}.docx"

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; filename="{filename}"'
            )
        },
    )

    if not feedback:
        return {
            "submitted": False,
            "feedback": None,
        }

    return {
        "submitted": True,
        "feedback": dict(feedback),
    }